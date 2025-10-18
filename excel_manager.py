"""
Excel file management for Telegram account credentials
"""
import os
from typing import List, Dict, Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config import EXCEL_FILE, EXCEL_HEADERS, REQUIRED_FIELDS


class ExcelManager:
    """Manages Excel file operations for account credentials"""
    
    def __init__(self, excel_file: str = str(EXCEL_FILE)):
        self.excel_file = excel_file
        self.headers = EXCEL_HEADERS
    
    def generate_template(self) -> bool:
        """Generate Excel template with headers and sample data"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Telegram Accounts"
            
            # Set headers
            for col, header in enumerate(self.headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # Add sample data
            sample_data = [
                ["+1234567890", "your_api_id_here", "your_api_hash_here"],
                ["+0987654321", "another_api_id", "another_api_hash"],
                ["", "", ""],  # Empty row for user to fill
                ["", "", ""],  # Another empty row
            ]
            
            for row, data in enumerate(sample_data, 2):
                for col, value in enumerate(data, 1):
                    ws.cell(row=row, column=col, value=value)
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(self.excel_file)
            print(f"✅ Excel template generated: {self.excel_file}")
            print("📝 Please fill in your account details and save the file.")
            return True
            
        except Exception as e:
            print(f"❌ Error generating Excel template: {e}")
            return False
    
    def load_accounts(self) -> List[Dict[str, str]]:
        """Load account credentials from Excel file"""
        if not os.path.exists(self.excel_file):
            print(f"❌ Excel file not found: {self.excel_file}")
            return []
        
        try:
            wb = load_workbook(self.excel_file)
            ws = wb.active
            
            accounts = []
            for row in range(2, ws.max_row + 1):  # Skip header row
                account = {}
                has_data = False
                
                for col, header in enumerate(self.headers, 1):
                    cell_value = ws.cell(row=row, column=col).value
                    if cell_value:
                        account[header] = str(cell_value).strip()
                        has_data = True
                
                # Only add accounts with data
                if has_data and self._validate_account(account):
                    accounts.append(account)
            
            print(f"✅ Loaded {len(accounts)} accounts from Excel file")
            return accounts
            
        except Exception as e:
            print(f"❌ Error loading Excel file: {e}")
            return []
    
    def _validate_account(self, account: Dict[str, str]) -> bool:
        """Validate account data"""
        for field in REQUIRED_FIELDS:
            if field not in account or not account[field]:
                print(f"⚠️ Skipping account with missing {field}")
                return False
        
        # Validate phone number format
        phone = account.get("Mobile_Number", "")
        # Auto-add + prefix if missing
        if not phone.startswith("+"):
            phone = "+" + phone
            account["Mobile_Number"] = phone
        
        if len(phone) < 10:
            print(f"⚠️ Skipping account with invalid phone number: {phone}")
            return False
        
        # Validate API key (should be numeric)
        try:
            int(account.get("API_Key", ""))
        except ValueError:
            print(f"⚠️ Skipping account with invalid API_Key: {account.get('API_Key')}")
            return False
        
        return True
    
    def add_account(self, mobile_number: str, api_key: str, hash_key: str) -> bool:
        """Add a new account to the Excel file"""
        try:
            if os.path.exists(self.excel_file):
                wb = load_workbook(self.excel_file)
                ws = wb.active
            else:
                # Create new workbook if file doesn't exist
                wb = Workbook()
                ws = wb.active
                ws.title = "Telegram Accounts"
                
                # Add headers
                for col, header in enumerate(self.headers, 1):
                    ws.cell(row=1, column=col, value=header)
            
            # Find next empty row
            next_row = ws.max_row + 1
            ws.cell(row=next_row, column=1, value=mobile_number)
            ws.cell(row=next_row, column=2, value=api_key)
            ws.cell(row=next_row, column=3, value=hash_key)
            
            wb.save(self.excel_file)
            print(f"✅ Account added: {mobile_number}")
            return True
            
        except Exception as e:
            print(f"❌ Error adding account: {e}")
            return False
    
    def get_account_count(self) -> int:
        """Get the number of accounts in the Excel file"""
        accounts = self.load_accounts()
        return len(accounts)
    
    def display_accounts(self) -> None:
        """Display all accounts in a formatted table"""
        accounts = self.load_accounts()
        
        if not accounts:
            print("📋 No accounts found in Excel file")
            return
        
        print(f"\n📋 Found {len(accounts)} accounts:")
        print("-" * 80)
        print(f"{'#':<3} {'Mobile Number':<15} {'API Key':<12} {'Hash Key':<20}")
        print("-" * 80)
        
        for i, account in enumerate(accounts, 1):
            mobile = account.get("Mobile_Number", "")
            api_key = account.get("API_Key", "")
            hash_key = account.get("Hash_Key", "")
            
            # Truncate long values for display
            display_hash = hash_key[:17] + "..." if len(hash_key) > 20 else hash_key
            display_api = api_key[:9] + "..." if len(api_key) > 12 else api_key
            
            print(f"{i:<3} {mobile:<15} {display_api:<12} {display_hash:<20}")
        
        print("-" * 80)


# Global instance
excel_manager = ExcelManager()
